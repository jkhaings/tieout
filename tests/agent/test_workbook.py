"""Tests for app.agent.workbook.write_commentary_sheet."""

from __future__ import annotations

from openpyxl import Workbook

from app.agent.workbook import write_commentary_sheet
from app.schemas import Chunk, Citation, Commentary


def _chunk(chunk_id: str, text: str, source_url: str = "https://example.com/filing.htm") -> Chunk:
    return Chunk(chunk_id=chunk_id, section="Item 7", text=text, source_url=source_url)


def test_refusal_renders_an_explicit_row_never_a_silent_skip() -> None:
    """Two production workbooks shipped with a Commentary tab holding only its
    headers, because every refusal was skipped. An absent row is not a
    fail-closed signal (CLAUDE.md rule 4) -- it is indistinguishable from a
    sheet that was never written."""
    wb = Workbook()
    commentary = [Commentary(line_item_key="revenue", text=None, citations=[])]
    write_commentary_sheet(
        wb,
        commentary,
        labels_by_key={"revenue": "Revenue"},
        chunks_by_id={},
        refusal_reasons={"revenue": "no filing passage was retrieved for this line item"},
    )

    ws = wb["Commentary"]
    data_rows = [row for row in ws.iter_rows(min_row=4, values_only=True) if any(row)]
    assert len(data_rows) == 1
    label, text, quote, source = data_rows[0]
    assert label == "Revenue"
    assert text == (
        "no grounded commentary available: no filing passage was retrieved for this line item"
    )
    assert quote is None and source is None


def test_refusal_without_a_recorded_reason_still_explains_itself() -> None:
    wb = Workbook()
    write_commentary_sheet(
        wb,
        [Commentary(line_item_key="revenue", text=None, citations=[])],
        labels_by_key={"revenue": "Revenue"},
        chunks_by_id={},
    )
    text = wb["Commentary"].cell(row=4, column=2).value
    assert str(text).startswith("no grounded commentary available: ")


def test_every_attempted_line_item_gets_a_row() -> None:
    """The invariant the audit would have caught: rows >= attempted items."""
    wb = Workbook()
    keys = ["revenue", "gross_profit", "net_income"]
    write_commentary_sheet(
        wb,
        [Commentary(line_item_key=k, text=None, citations=[]) for k in keys],
        labels_by_key={k: k.title() for k in keys},
        chunks_by_id={},
        refusal_reasons={"gross_profit": "this filer does not report this line item"},
    )
    data_rows = [row for row in wb["Commentary"].iter_rows(min_row=4, values_only=True) if any(row)]
    assert len(data_rows) == len(keys)
    assert all("no grounded commentary available" in str(row[1]) for row in data_rows)


def test_empty_commentary_says_narration_did_not_run() -> None:
    wb = Workbook()
    write_commentary_sheet(wb, [], labels_by_key={}, chunks_by_id={})
    data_rows = [row for row in wb["Commentary"].iter_rows(min_row=4, values_only=True) if any(row)]
    assert len(data_rows) == 1
    assert "narration did not run" in str(data_rows[0][1])


def test_writes_one_row_per_citation_with_source_resolved() -> None:
    wb = Workbook()
    chunk = _chunk("item7-0001", "Revenue grew due to strong iPhone sales.")
    commentary = [
        Commentary(
            line_item_key="revenue",
            text="Revenue grew, driven by iPhone sales.",
            citations=[Citation(chunk_id="item7-0001", quote="strong iPhone sales")],
        )
    ]
    write_commentary_sheet(
        wb, commentary, labels_by_key={"revenue": "Revenue"}, chunks_by_id={"item7-0001": chunk}
    )

    ws = wb["Commentary"]
    data_rows = [row for row in ws.iter_rows(min_row=4, values_only=True) if any(row)]
    assert len(data_rows) == 1
    label, text, quote, source = data_rows[0]
    assert label == "Revenue"
    assert text == "Revenue grew, driven by iPhone sales."
    assert quote == "strong iPhone sales"
    assert source == "https://example.com/filing.htm"


def test_multiple_citations_repeat_label_and_text_only_on_first_row() -> None:
    wb = Workbook()
    chunk_a = _chunk("item7-0001", "First point about revenue.")
    chunk_b = _chunk("item7-0002", "Second point about revenue.")
    commentary = [
        Commentary(
            line_item_key="revenue",
            text="Revenue commentary.",
            citations=[
                Citation(chunk_id="item7-0001", quote="First point"),
                Citation(chunk_id="item7-0002", quote="Second point"),
            ],
        )
    ]
    write_commentary_sheet(
        wb,
        commentary,
        labels_by_key={"revenue": "Revenue"},
        chunks_by_id={"item7-0001": chunk_a, "item7-0002": chunk_b},
    )

    ws = wb["Commentary"]
    data_rows = [row for row in ws.iter_rows(min_row=4, values_only=True) if any(row)]
    assert len(data_rows) == 2
    assert data_rows[0][0] == "Revenue" and data_rows[0][1] == "Revenue commentary."
    assert data_rows[1][0] is None and data_rows[1][1] is None
    assert data_rows[0][2] == "First point"
    assert data_rows[1][2] == "Second point"


def test_escapes_formula_injection_in_every_text_field() -> None:
    """SECURITY.md item 4: every text write must be escaped, no exceptions."""
    wb = Workbook()
    chunk = _chunk("item7-0001", "=EVIL(A1) is not a formula, just filing text.")
    commentary = [
        Commentary(
            line_item_key="revenue",
            text="=cmd|'/c calc'!A1",
            citations=[Citation(chunk_id="item7-0001", quote="=EVIL(A1) is not a formula")],
        )
    ]
    write_commentary_sheet(
        wb,
        commentary,
        labels_by_key={"revenue": "=Revenue"},
        chunks_by_id={"item7-0001": chunk},
    )

    ws = wb["Commentary"]
    label, text, quote, _source = next(
        row for row in ws.iter_rows(min_row=4, values_only=True) if any(row)
    )
    for value in (label, text, quote):
        assert value is not None
        assert value.startswith("'="), f"{value!r} was not escaped"


def test_missing_chunk_lookup_leaves_source_blank_without_raising() -> None:
    wb = Workbook()
    commentary = [
        Commentary(
            line_item_key="revenue",
            text="Some commentary.",
            citations=[Citation(chunk_id="unknown-0001", quote="quoted text here")],
        )
    ]
    write_commentary_sheet(wb, commentary, labels_by_key={"revenue": "Revenue"}, chunks_by_id={})

    ws = wb["Commentary"]
    row = next(r for r in ws.iter_rows(min_row=4, values_only=True) if any(r))
    assert row[3] is None
