"""Hermetic tests for app.model.excel -- workbooks are written under tmp_path, never
touching the real filesystem outside pytest's sandbox.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.edgar.tags import get as tags_get
from app.model.builder import share_filing_scale
from app.model.excel import (
    _NOT_AVAILABLE as NOT_AVAILABLE,
)
from app.model.excel import (
    _NOT_MEANINGFUL as NOT_MEANINGFUL,
)
from app.model.excel import build_workbook, escape_cell, write_workbook
from app.model.verifier import reconcile, verify
from app.schemas import LineItem, Statement, StatementSet, TieoutCheck, TieoutReport

_DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


# ---- escape_cell -----------------------------------------------------------------


@pytest.mark.parametrize("prefix", _DANGEROUS_PREFIXES)
def test_escape_cell_prefixes_dangerous_leading_characters(prefix: str) -> None:
    value = f'{prefix}HYPERLINK("https://evil.example/","clickme")'
    escaped = escape_cell(value)
    assert escaped == f"'{value}"


@pytest.mark.parametrize(
    "value", ["Apple Inc.", "Revenue", "Net income (loss)", "", "a=b", "100% grounded"]
)
def test_escape_cell_leaves_safe_strings_untouched(value: str) -> None:
    assert escape_cell(value) == value


def test_escaped_string_stays_a_string_cell_not_a_formula() -> None:
    """The actual mechanism escape_cell relies on: openpyxl converts an unescaped
    leading '=' into a live formula cell (data_type 'f'); escaping keeps it a
    plain string (data_type 's'), confirmed empirically before writing this test."""
    from openpyxl import Workbook

    hostile = '=HYPERLINK("https://evil.example/","clickme")'
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = hostile
    ws["A2"] = escape_cell(hostile)
    assert ws["A1"].data_type == "f"
    assert ws["A2"].data_type == "s"
    assert ws["A2"].value == "'" + hostile


# ---- build_workbook against real fixtures ------------------------------------------


def test_sheet_names_and_order(aapl_statements: StatementSet) -> None:
    report = verify(aapl_statements)
    recons = reconcile(aapl_statements)
    workbook = build_workbook(aapl_statements, report, recons)
    assert workbook.sheetnames == [
        "Income Statement",
        "Balance Sheet",
        "Cash Flow",
        "Ratios",
        "Tie-out",
    ]
    for name in workbook.sheetnames:
        assert len(name) <= 31  # Excel's sheet-name length limit
        assert not any(c in name for c in "[]:*?/\\")


def test_statement_sheets_have_one_row_per_matching_line_item(
    aapl_statements: StatementSet,
) -> None:
    report = verify(aapl_statements)
    recons = reconcile(aapl_statements)
    workbook = build_workbook(aapl_statements, report, recons)

    expected_income_items = sum(
        1 for item in aapl_statements.items if item.statement == Statement.INCOME
    )
    income_sheet = workbook["Income Statement"]
    label_column = [
        income_sheet.cell(row=r, column=1).value for r in range(4, income_sheet.max_row + 1)
    ]
    non_empty_labels = [v for v in label_column if v]
    assert len(non_empty_labels) == expected_income_items


def test_ratio_cells_are_formulas_not_pasted_numbers(aapl_statements: StatementSet) -> None:
    """Every ratio cell is a live formula or an explicit placeholder -- never a
    pasted number, and never blank (a blank reads as an oversight, and Excel
    treats it as zero in any arithmetic built on top of it)."""
    report = verify(aapl_statements)
    recons = reconcile(aapl_statements)
    workbook = build_workbook(aapl_statements, report, recons)
    ratios_sheet = workbook["Ratios"]

    found_formula = False
    for row in ratios_sheet.iter_rows(min_row=4):
        for cell in row:
            if cell.column == 1:
                continue
            assert isinstance(cell.value, str), (
                f"Ratios!{cell.coordinate} is {cell.value!r}, not a formula or placeholder"
            )
            assert cell.value.startswith("=") or cell.value in (NOT_AVAILABLE, NOT_MEANINGFUL), (
                f"Ratios!{cell.coordinate} is {cell.value!r}"
            )
            found_formula = found_formula or cell.value.startswith("=")
    assert found_formula


def test_gross_margin_formula_references_income_statement(aapl_statements: StatementSet) -> None:
    report = verify(aapl_statements)
    recons = reconcile(aapl_statements)
    workbook = build_workbook(aapl_statements, report, recons)
    ratios_sheet = workbook["Ratios"]

    gross_margin_row = next(
        r
        for r in range(4, ratios_sheet.max_row + 1)
        if ratios_sheet.cell(row=r, column=1).value == "Gross margin"
    )
    formula = ratios_sheet.cell(row=gross_margin_row, column=2).value
    assert isinstance(formula, str)
    assert formula.startswith("=")
    assert "'Income Statement'!" in formula


def test_first_year_revenue_growth_renders_not_available(aapl_statements: StatementSet) -> None:
    """There is no prior presented year for the first column to reference."""
    report = verify(aapl_statements)
    recons = reconcile(aapl_statements)
    workbook = build_workbook(aapl_statements, report, recons)
    ratios_sheet = workbook["Ratios"]

    growth_row = next(
        r
        for r in range(4, ratios_sheet.max_row + 1)
        if ratios_sheet.cell(row=r, column=1).value == "Revenue growth (YoY)"
    )
    first_year = ratios_sheet.cell(row=growth_row, column=2).value
    assert first_year == NOT_AVAILABLE
    assert not first_year.startswith("=")  # the invariant: no formula without a prior column
    assert str(ratios_sheet.cell(row=growth_row, column=3).value).startswith("=")


def test_tieout_sheet_lists_every_check_and_reconciliation(aapl_statements: StatementSet) -> None:
    report = verify(aapl_statements)
    recons = reconcile(aapl_statements)
    workbook = build_workbook(aapl_statements, report, recons)
    tieout_sheet = workbook["Tie-out"]

    all_values = [cell.value for row in tieout_sheet.iter_rows() for cell in row]
    for check in report.checks:
        assert check.check_id in all_values
    assert "PASS" in all_values  # every AAPL check passes
    assert "Reconciliations (not scored)" in all_values
    for recon in recons:
        assert recon.label in all_values


def test_failed_check_renders_fail_in_red(aapl_statements: StatementSet) -> None:
    broken_report = TieoutReport(
        ticker="AAPL",
        checks=[
            TieoutCheck(
                check_id="balance_sheet_equation_2024",
                description="Total assets = total liabilities + total stockholders' equity",
                fiscal_year=2024,
                passed=False,
                lhs=100.0,
                rhs=90.0,
                tolerance=1.0,
            )
        ],
    )
    workbook = build_workbook(aapl_statements, broken_report, [])
    tieout_sheet = workbook["Tie-out"]
    result_cell = next(
        cell for row in tieout_sheet.iter_rows() for cell in row if cell.value in ("PASS", "FAIL")
    )
    assert result_cell.value == "FAIL"
    assert result_cell.font.color.rgb.endswith("C0362C")


# ---- injection safety end-to-end, on a hostile synthetic StatementSet ----------------


def _hostile_statement_set() -> StatementSet:
    return StatementSet(
        ticker="EVIL",
        cik="0000000002",
        company_name='=HYPERLINK("https://evil.example/","clickme")',
        fiscal_years=[2024],
        items=[
            LineItem(
                key="total_assets",
                label="+cmd|' /C calc'!A1",
                statement=Statement.BALANCE,
                values={2024: 100.0},
            )
        ],
    )


def test_hostile_company_name_and_label_are_neutralized_end_to_end(tmp_path: Path) -> None:
    statements = _hostile_statement_set()
    report = TieoutReport(
        ticker="EVIL",
        checks=[
            TieoutCheck(
                check_id="=1+1",  # even a check_id is untrusted-shaped here
                description="@SUM(1,2)",
                fiscal_year=2024,
                passed=True,
            )
        ],
    )
    path = tmp_path / "evil.xlsx"
    write_workbook(statements, report, [], path)

    reloaded = load_workbook(path)

    balance_sheet = reloaded["Balance Sheet"]
    title_cell = balance_sheet.cell(row=1, column=1)
    assert title_cell.data_type == "s"
    assert isinstance(title_cell.value, str) and title_cell.value.startswith("'=HYPERLINK")

    label_cell = balance_sheet.cell(row=4, column=1)
    assert label_cell.data_type == "s"
    assert isinstance(label_cell.value, str) and label_cell.value.startswith("'+cmd")

    tieout_sheet = reloaded["Tie-out"]
    check_cells = [
        cell for row in tieout_sheet.iter_rows() for cell in row if cell.value == "'=1+1"
    ]
    assert check_cells, "escaped check_id not found as a plain string cell"
    assert check_cells[0].data_type == "s"


def test_write_workbook_creates_parent_directories(tmp_path: Path) -> None:
    statements = _hostile_statement_set()
    report = TieoutReport(ticker="EVIL", checks=[])
    nested_path = tmp_path / "nested" / "dir" / "model.xlsx"
    write_workbook(statements, report, [], nested_path)
    assert nested_path.exists()


# --- Production workbook audit: ratio cells never produce Excel errors -------
#
# Two real workbooks (META, MCD) shipped with ten #VALUE! ratio cells apiece,
# because a ratio formula was emitted whenever the *address* of an input cell
# existed -- including when that cell held the text "n/r".

_CELL_REF = re.compile(r"(?:'(?P<sheet>[^']+)'!)?\$?(?P<col>[A-Z]{1,3})\$?(?P<row>\d+)")


def _statements_with(values: dict[str, float], fy: int = 2024) -> StatementSet:
    """A single-year StatementSet carrying only the named line items."""
    return StatementSet(
        ticker="TEST",
        cik="0000000001",
        company_name="Test Co",
        fiscal_years=[fy],
        items=[
            LineItem(
                key=key,
                label=tags_get(key).label,
                statement=tags_get(key).statement,
                values={fy: value},
                unit=tags_get(key).unit,
            )
            for key, value in values.items()
        ],
    )


def _referenced_values(workbook: Workbook, sheet_title: str, formula: str) -> list[object]:
    """Resolve every cell reference in `formula` back to the value it points at."""
    out: list[object] = []
    for match in _CELL_REF.finditer(formula):
        target = workbook[match.group("sheet") or sheet_title]
        out.append(target[f"{match.group('col')}{match.group('row')}"].value)
    return out


@pytest.mark.parametrize("ticker", ["aapl", "msft", "mcd", "meta"])
def test_no_ratio_formula_references_a_non_numeric_cell(
    ticker: str, request: pytest.FixtureRequest
) -> None:
    """The acceptance criterion: zero error-producing formulas, proven structurally.

    A formula pointing at a text cell renders #VALUE! when Excel opens the
    workbook -- something no test can see by inspecting openpyxl values, so
    this resolves each reference by hand instead.
    """
    statements = request.getfixturevalue(f"{ticker}_statements")
    workbook = build_workbook(statements, verify(statements), reconcile(statements))
    ratios = workbook["Ratios"]

    checked = 0
    for row in ratios.iter_rows(min_row=4):
        for cell in row:
            if cell.column == 1 or not str(cell.value).startswith("="):
                continue
            for referenced in _referenced_values(workbook, "Ratios", str(cell.value)):
                # A same-sheet reference may point at another ratio's formula.
                if isinstance(referenced, str) and referenced.startswith("="):
                    continue
                assert isinstance(referenced, int | float), (
                    f"{statements.ticker} Ratios!{cell.coordinate} references "
                    f"{referenced!r}, which Excel would render as #VALUE!"
                )
                checked += 1
    assert checked, f"{statements.ticker} produced no ratio formulas to check"


@pytest.mark.parametrize("ticker", ["aapl", "msft", "mcd", "meta"])
def test_every_ratio_cell_is_a_formula_or_an_explicit_placeholder(
    ticker: str, request: pytest.FixtureRequest
) -> None:
    statements = request.getfixturevalue(f"{ticker}_statements")
    workbook = build_workbook(statements, verify(statements), reconcile(statements))
    for row in workbook["Ratios"].iter_rows(min_row=4):
        for cell in row:
            if cell.column == 1:
                continue
            assert isinstance(cell.value, str)
            assert cell.value.startswith("=") or cell.value in (NOT_AVAILABLE, NOT_MEANINGFUL)


def _ratio_row(workbook: Workbook, label: str) -> list[object]:
    ratios = workbook["Ratios"]
    row = next(
        r for r in range(4, ratios.max_row + 1) if ratios.cell(row=r, column=1).value == label
    )
    return [ratios.cell(row=row, column=c).value for c in range(2, ratios.max_column + 1)]


def test_mcd_gross_margin_is_not_available_without_gross_profit(
    mcd_statements: StatementSet,
) -> None:
    """McDonald's files no GrossProfit: five #VALUE! cells in the audited workbook."""
    workbook = build_workbook(mcd_statements, verify(mcd_statements), reconcile(mcd_statements))
    assert _ratio_row(workbook, "Gross margin") == [NOT_AVAILABLE] * 5


def test_meta_quick_ratio_is_not_available_without_inventory(
    meta_statements: StatementSet,
) -> None:
    """Meta reports no inventory, so the quick ratio has no subtrahend."""
    workbook = build_workbook(meta_statements, verify(meta_statements), reconcile(meta_statements))
    assert _ratio_row(workbook, "Quick ratio") == [NOT_AVAILABLE] * 5
    # The current ratio needs no inventory and must still compute.
    assert all(str(v).startswith("=") for v in _ratio_row(workbook, "Current ratio"))


def test_negative_equity_renders_not_meaningful(mcd_statements: StatementSet) -> None:
    """McDonald's equity is negative in all five presented years. A negative ROE
    on negative equity reads as a loss; the convention is to decline to show one."""
    equity = next(i for i in mcd_statements.items if i.key == "total_equity")
    assert all(v is not None and v < 0 for v in equity.values.values())

    workbook = build_workbook(mcd_statements, verify(mcd_statements), reconcile(mcd_statements))
    assert _ratio_row(workbook, "Return on equity") == [NOT_MEANINGFUL] * 5
    assert _ratio_row(workbook, "Debt-to-equity") == [NOT_MEANINGFUL] * 5
    # Ratios that do not divide by equity are unaffected.
    assert all(str(v).startswith("=") for v in _ratio_row(workbook, "Net margin"))


def test_zero_denominator_renders_not_meaningful() -> None:
    """A zero divisor would be #DIV/0!, which is an Excel error like any other."""
    statements = _statements_with({"revenue": 0.0, "gross_profit": 5.0})
    workbook = build_workbook(statements, verify(statements), reconcile(statements))
    assert _ratio_row(workbook, "Gross margin") == [NOT_MEANINGFUL]


def test_fcf_margin_inherits_the_fcf_placeholder() -> None:
    """FCF margin divides by the FCF row itself, so it must inherit that row's
    placeholder rather than emit a formula pointing at a text cell."""
    statements = _statements_with({"cfo": 100.0, "revenue": 500.0})  # no capex
    workbook = build_workbook(statements, verify(statements), reconcile(statements))
    assert _ratio_row(workbook, "Free cash flow") == [NOT_AVAILABLE]
    assert _ratio_row(workbook, "Free cash flow margin") == [NOT_AVAILABLE]


# --- Production workbook audit: share counts render on one scale -------------
#
# MCD filed weighted diluted shares as 751.8 (millions) while META filed
# 2,574,000,000 (units). Both are faithful to the filing; rendering them side
# by side without saying which is which is not.


def _share_row(workbook: Workbook, statements: StatementSet) -> tuple[str, list[object]]:
    ws = workbook["Income Statement"]
    row = next(
        r
        for r in range(4, ws.max_row + 1)
        if "diluted shares" in str(ws.cell(row=r, column=1).value)
    )
    return str(ws.cell(row=row, column=1).value), [
        ws.cell(row=row, column=c).value for c in range(2, 2 + len(statements.fiscal_years))
    ]


@pytest.mark.parametrize(
    ("ticker", "latest"),
    [("aapl", 15_004.7), ("msft", 7_453.0), ("mcd", 716.4), ("meta", 2_574.0)],
)
def test_share_counts_render_in_millions_whatever_the_filing_scale(
    ticker: str, latest: float, request: pytest.FixtureRequest
) -> None:
    statements = request.getfixturevalue(f"{ticker}_statements")
    workbook = build_workbook(statements, verify(statements), reconcile(statements))
    label, cells = _share_row(workbook, statements)
    assert label.endswith("(millions)")
    assert cells[-1] == pytest.approx(latest, abs=0.1)


def test_share_scale_prefers_the_eps_cross_check_over_magnitude() -> None:
    """A filer tagging share counts in thousands is invisible to a magnitude
    threshold but obvious from its own net income / EPS arithmetic."""
    statements = _statements_with(
        {"weighted_diluted_shares": 751_800.0, "eps_diluted": 11.39, "net_income": 8_223_000_000.0}
    )
    assert share_filing_scale(statements) == 1_000.0
    workbook = build_workbook(statements, verify(statements), reconcile(statements))
    assert _share_row(workbook, statements)[1] == [pytest.approx(751.8)]


def test_share_scale_falls_back_to_magnitude_without_eps() -> None:
    """No EPS evidence: thousands is deliberately unreachable, so a real count
    can never be misread by a factor of 1,000."""
    units = _statements_with({"weighted_diluted_shares": 2_574_000_000.0})
    millions = _statements_with({"weighted_diluted_shares": 751.8})
    assert share_filing_scale(units) == 1.0
    assert share_filing_scale(millions) == 1_000_000.0


def test_only_the_share_row_is_rescaled(mcd_statements: StatementSet) -> None:
    workbook = build_workbook(mcd_statements, verify(mcd_statements), reconcile(mcd_statements))
    ws = workbook["Income Statement"]
    by_label = {str(ws.cell(row=r, column=1).value): r for r in range(4, ws.max_row + 1)}
    last = 1 + len(mcd_statements.fiscal_years)
    assert ws.cell(row=by_label["Net income"], column=last).value == 8_563_000_000.0
    assert ws.cell(row=by_label["Diluted earnings per share"], column=last).value == 11.95


def test_display_scaling_never_mutates_the_model(mcd_statements: StatementSet) -> None:
    """`LineItem.values` stays exactly as filed -- `evals/tieout_eval.py` scores
    generated cells against gold values derived from the raw companyfacts."""
    build_workbook(mcd_statements, verify(mcd_statements), reconcile(mcd_statements))
    shares = next(i for i in mcd_statements.items if i.key == "weighted_diluted_shares")
    assert shares.values[2025] == 716.4


def test_derived_line_item_says_so_in_its_label(mcd_statements: StatementSet) -> None:
    """A value that follows from algebra rather than a filed tag is marked where
    it is read, not only in `LineItem.xbrl_tags`, which nothing renders."""
    workbook = build_workbook(mcd_statements, verify(mcd_statements), reconcile(mcd_statements))
    labels = [
        str(workbook["Balance Sheet"].cell(row=r, column=1).value)
        for r in range(4, workbook["Balance Sheet"].max_row + 1)
    ]
    assert any(label.startswith("Total liabilities") and "(derived)" in label for label in labels)
