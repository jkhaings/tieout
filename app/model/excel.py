"""Excel workbook generation: three statement tabs, live-formula ratios, a tie-out tab.

Every text write goes through `escape_cell()` (SECURITY.md item 4) --
including company names and check descriptions, since both ultimately
originate from a public SEC filing another party controls. Ratio cells are
always Excel formulas referencing statement cells
(`='Income Statement'!B5/'Income Statement'!B4`), never numbers computed in
Python and pasted in, so opening the workbook and clicking a ratio shows its
derivation, and the formula still recalculates if a cell is ever edited.

A line item with no reported value is written as the text `"n/r"`, not left
blank. Excel treats a genuinely blank cell as `0` inside arithmetic, which
would let a ratio formula silently compute a plausible-looking but wrong
number from missing data; text in the cell cannot be mistaken for zero.

The ratio writer therefore consults the model's own values before emitting
any formula, rather than pointing a formula at whatever is in a cell: a
ratio whose inputs are not all reported renders the text `"n/a"`, and one
whose inputs are present but whose result is undefined or misleading (a zero
divisor, or return on equity for a filer with negative equity) renders
`"n/m"`. Both are fail-closed in CLAUDE.md rule 4's sense, and legible --
an earlier version relied on `#VALUE!` propagating out of an `"n/r"` cell,
which is a visible failure but an unreadable one, and which a production
audit of two real workbooks found ten times apiece.
"""

from __future__ import annotations

from collections.abc import Callable, Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.edgar.tags import is_derived
from app.model.builder import share_filing_scale
from app.model.verifier import Reconciliation
from app.schemas import Statement, StatementSet, TieoutReport

_DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

_LABEL_COL = 1
_FIRST_YEAR_COL = 2
_HEADER_ROW = 3

_MONEY_FORMAT = "#,##0;(#,##0)"
_EPS_FORMAT = "#,##0.00;(#,##0.00)"
_PERCENT_FORMAT = "0.0%"
_RATIO_FORMAT = "0.00"
_SHARES_FORMAT = "#,##0.0;(#,##0.0)"

# A ratio cell is never blank and never an Excel error: it is a formula, or
# one of these two, which mean different things. "n/a" -- at least one input
# was not reported, so no formula can be written. "n/m" -- every input is
# present but the result would not be meaningful: a zero divisor, or a ratio
# on non-positive equity, where the finance convention is to decline to show
# a number rather than print a misleading negative one.
_NOT_AVAILABLE = "n/a"
_NOT_MEANINGFUL = "n/m"

# Share counts are filed at whatever scale the filer presents; the workbook
# renders them all in millions and says so in the row label, so two companies
# can be read side by side. See `app.model.builder.share_filing_scale`.
_SHARE_DISPLAY_DIVISOR = 1_000_000.0
_SHARES_LABEL_SUFFIX = " (millions)"

CellRegistry = dict[tuple[str, int], str]


def escape_cell(value: str) -> str:
    """Neutralize formula injection: prefix a leading dangerous character with `'`.

    Excel (and every major spreadsheet app) treats a cell beginning with
    `=`, `+`, `-`, `@`, a tab, or a carriage return as a formula. A leading
    `'` forces text interpretation without changing what a user sees. Apply
    to every string written to a cell, no exceptions (SECURITY.md item 4).
    """
    if value and value[0] in _DANGEROUS_PREFIXES:
        return "'" + value
    return value


def _quote_sheet(title: str) -> str:
    """Quote a sheet name for use in a formula reference, if it needs it."""
    return f"'{title}'" if any(c in title for c in " -") else title


def _plain_cell(ws: Worksheet, row: int, column: int) -> Cell:
    """`ws.cell(row, column)` narrowed from `Cell | MergedCell` to `Cell`.

    This workbook never calls `merge_cells`, so a `MergedCell` -- which the
    stub return type allows for in general -- can never actually occur
    here; asserting that lets callers set `.value`/`.font`/`.number_format`
    without every such call site repeating a `# type: ignore`.
    """
    cell = ws.cell(row=row, column=column)
    assert isinstance(cell, Cell)
    return cell


def _write_statement_sheet(
    ws: Worksheet,
    statements: StatementSet,
    statement: Statement,
    registry: CellRegistry,
) -> None:
    """Write one statement tab (Income/Balance/Cash Flow) and register every cell address."""
    ws.cell(
        row=1, column=1, value=escape_cell(f"{statements.company_name} ({statements.ticker})")
    ).font = Font(bold=True, size=13)
    ws.cell(
        row=2,
        column=1,
        value=("All amounts in U.S. dollars except per-share data; share counts in millions."),
    ).font = Font(italic=True, size=9)

    ws.cell(row=_HEADER_ROW, column=_LABEL_COL, value="Line item").font = Font(bold=True)
    for i, fy in enumerate(statements.fiscal_years):
        header_cell = ws.cell(row=_HEADER_ROW, column=_FIRST_YEAR_COL + i, value=f"FY{fy}")
        header_cell.font = Font(bold=True)
        header_cell.alignment = Alignment(horizontal="right")

    row = _HEADER_ROW + 1
    quoted_title = _quote_sheet(ws.title)
    share_scale = share_filing_scale(statements)
    for item in statements.items:
        if item.statement != statement:
            continue
        in_shares = item.unit == "shares"
        label = item.label
        if in_shares:
            label += _SHARES_LABEL_SUFFIX
        if any(is_derived(tag) for tag in item.xbrl_tags):
            # This filer never filed the concept; it follows exactly from two
            # others (`app.model.builder`). Say so where the number is read.
            label += " (derived)"
        ws.cell(row=row, column=_LABEL_COL, value=escape_cell(label))
        if in_shares:
            number_format = _SHARES_FORMAT
        elif item.key == "eps_diluted":
            number_format = _EPS_FORMAT
        else:
            number_format = _MONEY_FORMAT
        for i, fy in enumerate(statements.fiscal_years):
            col = _FIRST_YEAR_COL + i
            cell = _plain_cell(ws, row, col)
            value = item.values.get(fy)
            if value is None:
                cell.value = escape_cell("n/r")
                cell.font = Font(italic=True, color="999999")
                cell.alignment = Alignment(horizontal="right")
            else:
                # Display-only rescale: `LineItem.values` stays exactly as
                # filed, which is what `evals/tieout_eval.py` scores against.
                cell.value = value * share_scale / _SHARE_DISPLAY_DIVISOR if in_shares else value
                cell.number_format = number_format
            registry[(item.key, fy)] = f"{quoted_title}!{get_column_letter(col)}{row}"
        row += 1


def _write_ratios_sheet(ws: Worksheet, statements: StatementSet, registry: CellRegistry) -> None:
    """Write the Ratios tab. Every cell is a live formula, `"n/a"`, or `"n/m"` -- never blank.

    The registry says *where* a statement cell is; it says nothing about what
    is in it, and `_write_statement_sheet` registers an address for an
    unreported value's `"n/r"` text cell just as it does for a number. So the
    availability decision is made here against the model's own values, before
    any formula is emitted.
    """
    ws.cell(
        row=1,
        column=1,
        value=escape_cell(f"{statements.company_name} ({statements.ticker}) — Ratios"),
    ).font = Font(bold=True, size=13)
    ws.cell(
        row=2,
        column=1,
        value=(
            "Every populated cell is a live formula referencing the statement tabs. "
            '"n/a" = an input is not reported; "n/m" = not meaningful.'
        ),
    ).font = Font(italic=True, size=9)

    ws.cell(row=_HEADER_ROW, column=_LABEL_COL, value="Ratio").font = Font(bold=True)
    for i, fy in enumerate(statements.fiscal_years):
        ws.cell(row=_HEADER_ROW, column=_FIRST_YEAR_COL + i, value=f"FY{fy}").font = Font(bold=True)

    row = _HEADER_ROW + 1
    # Same-sheet refs, for ratios built from other ratio rows (e.g. FCF margin),
    # and the placeholder each such row rendered instead of a formula, if any.
    own_cells: CellRegistry = {}
    own_flags: dict[tuple[str, int], str] = {}

    values_by_key = {item.key: item.values for item in statements.items}
    # Cells whose displayed number is rescaled for reading (share counts, in
    # millions). A formula must never reference one: the cell no longer holds
    # the as-filed value. No ratio below uses one; this keeps it that way.
    scaled_keys = {item.key for item in statements.items if item.unit == "shares"}

    def value(key: str, fy: int) -> float | None:
        """The model value behind the statement cell for `key`/`fy`, if reported."""
        return values_by_key.get(key, {}).get(fy)

    def numeric(key: str, fy: int) -> bool:
        """True when the statement tab holds a number here, not `"n/r"` or nothing."""
        return value(key, fy) is not None and (key, fy) in registry

    def ref(key: str, fy: int) -> str:
        """The statement-tab cell address for `key`/`fy`. Only call once `numeric` holds."""
        assert key not in scaled_keys, f"{key} is display-scaled; a formula must not reference it"
        return registry[(key, fy)]

    def ratio_cell(
        fy: int,
        *,
        inputs: Sequence[tuple[str, int]],
        build: Callable[[], str],
        nonzero: Sequence[tuple[str, int]] = (),
        positive: Sequence[tuple[str, int]] = (),
    ) -> str:
        """A formula, or the text that replaces it. Availability is decided first.

        Args:
            fy: The fiscal year this cell covers (for the error text only).
            inputs: Every `(key, fiscal_year)` the formula would reference.
            build: Builds the formula, called only once every input is numeric.
            nonzero: Inputs that would be divided by, so zero means `#DIV/0!`.
            positive: Inputs that must be strictly positive for the result to
                mean anything -- equity, for return on equity and
                debt-to-equity, where a negative denominator flips the sign of
                a ratio readers interpret as a level.
        """
        del fy  # signature symmetry with the per-year callables below
        if any(not numeric(key, year) for key, year in inputs):
            return _NOT_AVAILABLE
        if any(value(key, year) == 0 for key, year in nonzero):
            return _NOT_MEANINGFUL
        if any((value(key, year) or 0.0) <= 0 for key, year in positive):
            return _NOT_MEANINGFUL
        return build()

    def write_row(
        label: str, key: str, number_format: str, cell_for_year: Callable[[int], str]
    ) -> None:
        """Write one ratio row, and register both its cells and its placeholders."""
        nonlocal row
        ws.cell(row=row, column=_LABEL_COL, value=escape_cell(label))
        for i, fy in enumerate(statements.fiscal_years):
            col = _FIRST_YEAR_COL + i
            rendered = cell_for_year(fy)
            cell = _plain_cell(ws, row, col)
            if rendered.startswith("="):
                cell.value = rendered
                cell.number_format = number_format
            else:
                cell.value = escape_cell(rendered)
                cell.font = Font(italic=True, color="999999")
                cell.alignment = Alignment(horizontal="right")
                own_flags[(key, fy)] = rendered
            own_cells[(key, fy)] = f"{get_column_letter(col)}{row}"
        row += 1

    def simple_ratio(
        numerator: str, denominator: str, *, positive_denominator: bool = False
    ) -> Callable[[int], str]:
        """A `=numerator/denominator` row over two statement line items."""

        def render(fy: int) -> str:
            return ratio_cell(
                fy,
                inputs=((numerator, fy), (denominator, fy)),
                nonzero=() if positive_denominator else ((denominator, fy),),
                positive=((denominator, fy),) if positive_denominator else (),
                build=lambda: f"={ref(numerator, fy)}/{ref(denominator, fy)}",
            )

        return render

    def quick_ratio(fy: int) -> str:
        """`(current assets - inventory) / current liabilities`."""
        return ratio_cell(
            fy,
            inputs=(
                ("total_current_assets", fy),
                ("inventory", fy),
                ("total_current_liabilities", fy),
            ),
            nonzero=(("total_current_liabilities", fy),),
            build=lambda: (
                f"=({ref('total_current_assets', fy)}-{ref('inventory', fy)})"
                f"/{ref('total_current_liabilities', fy)}"
            ),
        )

    def fcf(fy: int) -> str:
        """Free cash flow: `CFO - capital expenditures`."""
        return ratio_cell(
            fy,
            inputs=(("cfo", fy), ("capex", fy)),
            build=lambda: f"={ref('cfo', fy)}-{ref('capex', fy)}",
        )

    def fcf_margin(fy: int) -> str:
        """Free cash flow as a fraction of revenue; references this sheet's own FCF row.

        Inherits the FCF row's placeholder rather than deciding again: a
        formula dividing by a cell that itself reads `"n/a"` is exactly the
        defect this module stopped emitting.
        """
        assert ("fcf", fy) in own_cells, "the FCF row must be written before FCF margin"
        inherited = own_flags.get(("fcf", fy))
        if inherited is not None:
            return inherited
        return ratio_cell(
            fy,
            inputs=(("revenue", fy),),
            nonzero=(("revenue", fy),),
            build=lambda: f"={own_cells[('fcf', fy)]}/{ref('revenue', fy)}",
        )

    def yoy_revenue_growth(fy: int) -> str:
        """Year-over-year revenue growth; `"n/a"` for the first presented year."""
        index = statements.fiscal_years.index(fy)
        if index == 0:
            return _NOT_AVAILABLE  # no prior column on this sheet to reference
        prior_fy = statements.fiscal_years[index - 1]
        return ratio_cell(
            fy,
            inputs=(("revenue", fy), ("revenue", prior_fy)),
            nonzero=(("revenue", prior_fy),),
            build=lambda: (
                f"=({ref('revenue', fy)}-{ref('revenue', prior_fy)})/{ref('revenue', prior_fy)}"
            ),
        )

    write_row(
        "Gross margin", "gross_margin", _PERCENT_FORMAT, simple_ratio("gross_profit", "revenue")
    )
    write_row(
        "Operating margin",
        "operating_margin",
        _PERCENT_FORMAT,
        simple_ratio("operating_income", "revenue"),
    )
    write_row("Net margin", "net_margin", _PERCENT_FORMAT, simple_ratio("net_income", "revenue"))
    write_row(
        "Return on assets", "roa", _PERCENT_FORMAT, simple_ratio("net_income", "total_assets")
    )
    write_row(
        "Return on equity",
        "roe",
        _PERCENT_FORMAT,
        simple_ratio("net_income", "total_equity", positive_denominator=True),
    )
    write_row(
        "Current ratio",
        "current_ratio",
        _RATIO_FORMAT,
        simple_ratio("total_current_assets", "total_current_liabilities"),
    )
    write_row("Quick ratio", "quick_ratio", _RATIO_FORMAT, quick_ratio)
    write_row(
        "Asset turnover", "asset_turnover", _RATIO_FORMAT, simple_ratio("revenue", "total_assets")
    )
    write_row(
        "Debt-to-equity",
        "debt_to_equity",
        _RATIO_FORMAT,
        simple_ratio("total_liabilities", "total_equity", positive_denominator=True),
    )
    write_row("Free cash flow", "fcf", _MONEY_FORMAT, fcf)
    write_row("Free cash flow margin", "fcf_margin", _PERCENT_FORMAT, fcf_margin)
    write_row("Revenue growth (YoY)", "revenue_growth", _PERCENT_FORMAT, yoy_revenue_growth)


def _write_tieout_sheet(
    ws: Worksheet,
    statements: StatementSet,
    tieout: TieoutReport,
    reconciliations: list[Reconciliation],
) -> None:
    """Write the Tie-out tab: scored checks, then informational reconciliations."""
    ws.cell(
        row=1,
        column=1,
        value=escape_cell(f"{statements.company_name} ({statements.ticker}) — Tie-out"),
    ).font = Font(bold=True, size=13)

    row = 3
    ws.cell(row=row, column=1, value="Scored checks").font = Font(bold=True, size=11)
    row += 1
    for col, header in enumerate(
        [
            "Check",
            "Fiscal year",
            "Description",
            "Left side",
            "Right side",
            "Difference",
            "Tolerance",
            "Result",
        ],
        start=1,
    ):
        ws.cell(row=row, column=col, value=header).font = Font(bold=True)
    row += 1
    for check in tieout.checks:
        diff = None if check.lhs is None or check.rhs is None else check.lhs - check.rhs
        ws.cell(row=row, column=1, value=escape_cell(check.check_id))
        ws.cell(row=row, column=2, value=check.fiscal_year)
        ws.cell(row=row, column=3, value=escape_cell(check.description))
        if check.lhs is not None:
            ws.cell(row=row, column=4, value=check.lhs).number_format = _MONEY_FORMAT
        if check.rhs is not None:
            ws.cell(row=row, column=5, value=check.rhs).number_format = _MONEY_FORMAT
        if diff is not None:
            ws.cell(row=row, column=6, value=diff).number_format = _MONEY_FORMAT
        tolerance_cell = ws.cell(row=row, column=7, value=check.tolerance)
        tolerance_cell.number_format = _MONEY_FORMAT
        result_cell = ws.cell(row=row, column=8, value="PASS" if check.passed else "FAIL")
        result_cell.font = Font(bold=True, color="1A7F37" if check.passed else "C0362C")
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Reconciliations (not scored)").font = Font(bold=True, size=11)
    row += 1
    note = (
        "Retained-earnings roll-forward does not tie exactly from XBRL data alone "
        "(buybacks charged to retained earnings are not consistently tagged); shown "
        "here as a named residual, not a pass/fail check -- see app.model.verifier.reconcile."
    )
    ws.cell(row=row, column=1, value=escape_cell(note)).font = Font(italic=True, size=9)
    row += 1
    for col, header in enumerate(
        [
            "Item",
            "Fiscal year",
            "Beginning RE",
            "Net income",
            "Dividends",
            "Buybacks",
            "Ending RE (actual)",
            "Ending RE (expected)",
            "Residual",
            "Residual % of assets",
        ],
        start=1,
    ):
        ws.cell(row=row, column=col, value=header).font = Font(bold=True)
    row += 1
    for recon in reconciliations:
        ws.cell(row=row, column=1, value=escape_cell(recon.label))
        ws.cell(row=row, column=2, value=recon.fiscal_year)
        money_values = (
            recon.beginning,
            recon.net_income,
            recon.dividends,
            recon.buybacks,
            recon.ending_actual,
            recon.ending_expected,
            recon.residual,
        )
        for col, value in enumerate(money_values, start=3):
            ws.cell(row=row, column=col, value=value).number_format = _MONEY_FORMAT
        if recon.residual_pct_of_assets is not None:
            ws.cell(
                row=row, column=10, value=recon.residual_pct_of_assets / 100
            ).number_format = _PERCENT_FORMAT
        row += 1


def build_workbook(
    statements: StatementSet,
    tieout: TieoutReport,
    reconciliations: list[Reconciliation],
) -> Workbook:
    """Build the full workbook in memory: three statement tabs, Ratios, Tie-out.

    Pure and I/O-free so tests can inspect the result directly; `write_workbook`
    below is the thin convenience wrapper that saves it to disk.
    """
    workbook = Workbook()
    registry: CellRegistry = {}

    income_sheet = workbook.active
    assert income_sheet is not None
    income_sheet.title = "Income Statement"
    _write_statement_sheet(income_sheet, statements, Statement.INCOME, registry)

    balance_sheet = workbook.create_sheet("Balance Sheet")
    _write_statement_sheet(balance_sheet, statements, Statement.BALANCE, registry)

    cashflow_sheet = workbook.create_sheet("Cash Flow")
    _write_statement_sheet(cashflow_sheet, statements, Statement.CASHFLOW, registry)

    ratios_sheet = workbook.create_sheet("Ratios")
    _write_ratios_sheet(ratios_sheet, statements, registry)

    tieout_sheet = workbook.create_sheet("Tie-out")
    _write_tieout_sheet(tieout_sheet, statements, tieout, reconciliations)

    for sheet in (income_sheet, balance_sheet, cashflow_sheet, ratios_sheet, tieout_sheet):
        sheet.column_dimensions["A"].width = 44
        for i in range(len(statements.fiscal_years)):
            sheet.column_dimensions[get_column_letter(_FIRST_YEAR_COL + i)].width = 16

    return workbook


def write_workbook(
    statements: StatementSet,
    tieout: TieoutReport,
    reconciliations: list[Reconciliation],
    path: Path,
) -> None:
    """Build the workbook and save it to `path`, creating parent directories as needed."""
    workbook = build_workbook(statements, tieout, reconciliations)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
